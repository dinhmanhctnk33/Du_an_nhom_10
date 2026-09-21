import sys
import docx
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

src_file = Path('02_GenAI_SoftwareDevelopment_requirements-qa.docx')
out_file = Path('02_fixed.docx')

print(f'Loading {src_file}...')
doc = docx.Document(src_file)

# 1. Update Table 8 (AI Requirements - AIR)
t8 = doc.tables[8]
print('Original Table 8:')
for row in t8.rows:
    print([c.text.strip() for c in row.cells])

# Table 8 content fix
air_data = [
    ('AIR-DRAFT-001', 'Tích hợp AI sinh báo cáo lưu lượng xe, hỏi đáp dữ liệu và gợi ý nhân sự'),
    ('AIR-DRAFT-002', 'Quản lý luồng dữ liệu đầu vào (lượt xe, doanh thu, lấp đầy, khung giờ) và đầu ra AI'),
    ('AIR-DRAFT-003', 'Thiết lập prompt có ràng buộc chống tự tạo số liệu (chỉ nhận xét từ dữ liệu được cung cấp)')
]

for idx, (code, text) in enumerate(air_data, start=1):
    row = t8.rows[idx]
    # Set cell 0
    p0 = row.cells[0].paragraphs[0]
    p0.clear()
    r0 = p0.add_run(code)
    r0.font.name = 'Calibri'
    r0.font.size = docx.shared.Pt(10)
    
    # Set cell 1
    p1 = row.cells[1].paragraphs[0]
    p1.clear()
    r1 = p1.add_run(text)
    r1.font.name = 'Calibri'
    r1.font.size = docx.shared.Pt(10)

print('\nUpdated Table 8:')
for row in t8.rows:
    print([c.text.strip() for c in row.cells])

# 2. Update P007 reference string if needed for explicit file name match
p7 = doc.paragraphs[7]
if 'Project Plan' in p7.text and '01_GenAI_SoftwareDevelopment_project-plan.docx' not in p7.text:
    p7.text = p7.text.replace('Project Plan đã hoàn thiện', '01_GenAI_SoftwareDevelopment_project-plan.docx')

# Save to 02_fixed.docx
doc.save(out_file)
print('====================================================')
print(f'SUCCESS! File saved to: {out_file.resolve()}')
print('====================================================')
