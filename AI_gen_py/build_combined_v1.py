import sys
import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

# Import data structures from existing scripts
print("Importing use cases and activity cases...")
from build_sequence_v2 import USE_CASES
from build_activity_v1 import ACTIVITY_CASES
print("Import successful!")

# ---------------------------------------------------------
# Helper functions for python-docx styling & formatting
# ---------------------------------------------------------
def fmt(run, size=10, bold=None, italic=None, color=None, font_name='Calibri'):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rPr.append(rFonts)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_cell(cell, text, bold=False, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    fmt(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def set_table_col_widths(table, widths_in_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in_inches):
            cell.width = Inches(w)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    set_table_col_widths(tbl, [6.5])
    c = tbl.cell(0, 0)
    shade_cell(c, 'F8FAFC')
    p = c.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(code_str)
    fmt(r, size=7.5, bold=False, color=(30, 41, 59), font_name='Consolas')
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def generate_plantuml(title, code, participants, steps):
    out = [
        '@startuml',
        f'title {code} — {title}',
        'hide footbox',
        'skinparam sequenceMessageAlign center',
        'skinparam responseMessageBelowArrow true'
    ]
    aliases = {}
    for i, p in enumerate(participants):
        aliases[p] = f'P{i+1}'
        out.append(f'participant "{p}" as P{i+1}')
        
    curr_branch = None
    for branch, src, dst, msg, is_return in steps:
        if branch != curr_branch:
            if curr_branch is None:
                out.append(f'alt {branch}')
            else:
                out.append(f'else {branch}')
            curr_branch = branch
            
        arrow = '--->' if is_return else '->'
        out.append(f'{aliases[src]} {arrow} {aliases[dst]} : {msg}')
        
    out.append('end')
    out.append('@enduml')
    return '\n'.join(out)

def generate_activity_plantuml(title, idx_num, partitions, steps_def):
    out = [
        '@startuml',
        f'title {idx_num}. Sơ đồ hoạt động của {title}',
        'skinparam ActivityFontSize 13',
        'skinparam ActivityBorderColor #1E293B',
        'skinparam ActivityBackgroundColor #F8FAFC',
        'skinparam ActivityDiamondBorderColor #D97706',
        'skinparam ActivityDiamondBackgroundColor #FEF3C7'
    ]
    
    curr_part = None
    for part, action_type, text, cond in steps_def:
        if part != curr_part:
            out.append(f'|{part}|')
            curr_part = part
            
        if action_type == 'start':
            out.append('start')
        elif action_type == 'stop':
            out.append('stop')
        elif action_type == 'action':
            out.append(f':{text};')
        elif action_type == 'decision':
            out.append(f'if ({text}?) then ([Đúng])')
        elif action_type == 'else':
            out.append(f'else ([{cond}])')
            out.append(f'  :{text};')
        elif action_type == 'endif':
            out.append('endif')
            
    out.append('@enduml')
    return '\n'.join(out)

# ---------------------------------------------------------
# Build Sơ_đồ_tuần_tự_hoạt_động.docx
# ---------------------------------------------------------
print("Starting assembly of Sơ_đồ_tuần_tự_hoạt_động.docx...")
doc_target = Document()

# Set standard margins (1 inch = 72 pt)
for s in doc_target.sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

# Document Title Block
p_title = doc_target.add_paragraph()
r_title = p_title.add_run('HỆ THỐNG QUẢN LÝ BÃI XE THÔNG MINH')
fmt(r_title, size=18, bold=True, color=(15, 23, 42))

p_sub = doc_target.add_paragraph()
r_sub = p_sub.add_run('TÀI LIỆU TỔNG HỢP SƠ ĐỒ TUẦN TỰ & SƠ ĐỒ HOẠT ĐỘNG')
fmt(r_sub, size=14, bold=True, color=(37, 99, 235))

p_sub2 = doc_target.add_paragraph()
r_sub2 = p_sub2.add_run('Đặc tả Sơ đồ tuần tự (UML 2.5 v2.0) và Sơ đồ hoạt động (UML 2.5 v1.0) cho 12 Use Case tiêu chuẩn')
fmt(r_sub2, size=10.5, italic=True, color=(71, 85, 105))

# Document Metadata Table
tbl_meta = doc_target.add_table(rows=4, cols=2)
set_table_col_widths(tbl_meta, [2.0, 4.5])
meta_data = [
    ('Nguồn đặc tả', 'Đặc tả UC_done.docx (Nhóm 10 — Đinh Tiến Mạnh, Tạ Văn Đức)'),
    ('Phạm vi hệ thống', '12 Use Case tiêu chuẩn (UC-001 đến UC-012)'),
    ('Quy ước sơ đồ', 'Sequence Diagram (alt/else/except) & Activity Diagram (Swimlanes/Decisions)'),
    ('Phiên bản tài liệu', 'v1.0 — Tích hợp đồng bộ Sơ đồ tuần tự và Sơ đồ hoạt động kèm Mã nguồn StarUML')
]
for i, (k, v) in enumerate(meta_data):
    row = tbl_meta.rows[i]
    set_cell(row.cells[0], k, bold=True, size=9, color=(15, 23, 42))
    shade_cell(row.cells[0], 'F1F5F9')
    set_cell(row.cells[1], v, bold=False, size=9)

p_space = doc_target.add_paragraph()
p_space.paragraph_format.space_after = Pt(12)

# Section 1: Conventions & Overview
h1_1 = doc_target.add_heading(level=1)
r = h1_1.add_run('Hướng dẫn & Quy ước Thiết kế Sơ đồ')
fmt(r, size=14, bold=True, color=(15, 23, 42))

p_conv = doc_target.add_paragraph()
p_conv.paragraph_format.line_spacing = 1.15
r = p_conv.add_run(
    'Tài liệu này tổng hợp toàn bộ các sơ đồ thiết kế cho 12 Use Case của hệ thống, bao gồm:\n'
    '1. Sơ đồ tuần tự (Sequence Diagram): Biểu diễn các thông điệp truyền tải giữa Actor, Giao diện, Dịch vụ nghiệp vụ và Kho dữ liệu. '
    'Luồng chính được đặt trong khối alt; các luồng thay thế (A) và ngoại lệ (E) được đặt trong các khối else tương ứng.\n'
    '2. Sơ đồ hoạt động (Activity Diagram): Biểu diễn luồng tiến trình công việc phân theo làn trách nhiệm (Swimlanes/Partition). '
    'Các điểm quyết định (Decision Node) biểu diễn phân nhánh luồng nghiệp vụ tương ứng.\n'
    '3. Mã nguồn vẽ sơ đồ: Cung cấp đầy đủ mã PlantUML cho cả hai loại sơ đồ, sẵn sàng nạp trực tiếp vào StarUML hoặc các trình biên dịch PlantUML.'
)
fmt(r, size=10, color=(30, 41, 59))

# Section 2: Detailed 12 Use Cases
for i in range(12):
    seq_case = USE_CASES[i]
    act_case = ACTIVITY_CASES[i]
    idx = i + 1
    title = seq_case['title']
    uc_id = seq_case['uc_id']
    
    heading_text = f"{idx}. Sơ đồ tuần tự và hoạt động của {title}"
    print(f"[{idx}/12] Adding combined content for {heading_text}...")
    doc_target.add_page_break()
    
    # Heading 1
    h_uc = doc_target.add_heading(level=1)
    r = h_uc.add_run(heading_text)
    fmt(r, size=13, bold=True, color=(15, 23, 42))
    
    # Purpose & Overview Paragraph
    p_purpose = doc_target.add_paragraph()
    p_purpose.paragraph_format.space_after = Pt(6)
    r = p_purpose.add_run(f"Mã Use Case: {uc_id} | {seq_case['purpose']}\n{act_case['overview']}")
    fmt(r, size=9.5, italic=True, color=(51, 65, 85))
    
    # ---------------------------------------------------------
    # Subsection 1: Sequence Diagram
    # ---------------------------------------------------------
    h2_seq = doc_target.add_heading(level=2)
    r = h2_seq.add_run(f"{idx}.1. Sơ đồ tuần tự của {title}")
    fmt(r, size=11, bold=True, color=(30, 41, 59))
    
    # Insert Sequence Image
    seq_img_name = f"{seq_case['code']}_v2.png"
    seq_img_path = Path('v2_assets') / seq_img_name
    if seq_img_path.exists():
        doc_target.add_picture(str(seq_img_path), width=Inches(6.5))
        p_cap = doc_target.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(f"Hình — {seq_case['code']}: Sơ đồ tuần tự {uc_id} — {title}")
        fmt(r_cap, size=8.5, italic=True, color=(100, 116, 139))
    else:
        print(f"Warning: Sequence image {seq_img_path} not found!")

    # Message Table
    p_tbl_seq_lbl = doc_target.add_paragraph()
    p_tbl_seq_lbl.paragraph_format.space_after = Pt(4)
    r = p_tbl_seq_lbl.add_run("Bảng chuỗi thông điệp chi tiết:")
    fmt(r, size=9.5, bold=True, color=(30, 41, 59))
    
    tbl_msg = doc_target.add_table(rows=1, cols=5)
    set_table_col_widths(tbl_msg, [0.45, 1.85, 1.1, 1.1, 2.0])
    msg_headers = ['#', 'Phân nhánh / Điều kiện', 'Từ', 'Đến', 'Thông điệp / Hành vi']
    for c, h in zip(tbl_msg.rows[0].cells, msg_headers):
        set_cell(c, h, bold=True, size=8.5, color=(255, 255, 255))
        shade_cell(c, '334155')
        
    for st_idx, (branch, src, dst, msg, is_ret) in enumerate(seq_case['steps'], start=1):
        row = tbl_msg.add_row()
        set_cell(row.cells[0], str(st_idx), bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        
        set_cell(row.cells[1], branch, bold=True, size=7.5)
        if 'Main' in branch or 'Chính' in branch:
            shade_cell(row.cells[1], 'F0FDF4')
        elif 'E' in branch or 'Ngoại lệ' in branch:
            shade_cell(row.cells[1], 'FEF2F2')
        else:
            shade_cell(row.cells[1], 'FFFBEB')
            
        set_cell(row.cells[2], src, bold=False, size=8)
        set_cell(row.cells[3], dst, bold=False, size=8)
        
        msg_fmt = f"{msg} (Return)" if is_ret else msg
        set_cell(row.cells[4], msg_fmt, bold=False, size=8)
        
    p_sp1 = doc_target.add_paragraph()
    p_sp1.paragraph_format.space_after = Pt(6)
    
    # PlantUML Sequence Code
    h3_seq_code = doc_target.add_heading(level=2)
    r = h3_seq_code.add_run(f"Mã PlantUML sơ đồ tuần tự của {title}")
    fmt(r, size=11, bold=True, color=(30, 41, 59))
    
    puml_seq_str = generate_plantuml(seq_case['title'], seq_case['code'], seq_case['participants'], seq_case['steps'])
    add_code_block(doc_target, puml_seq_str)

    # ---------------------------------------------------------
    # Subsection 2: Activity Diagram
    # ---------------------------------------------------------
    h2_act = doc_target.add_heading(level=2)
    r = h2_act.add_run(f"{idx}.2. Sơ đồ hoạt động của {title}")
    fmt(r, size=11, bold=True, color=(30, 41, 59))
    
    # Insert Activity Image
    act_img_name = f"act_{idx:02d}.png"
    act_img_path = Path('act_assets') / act_img_name
    if act_img_path.exists():
        doc_target.add_picture(str(act_img_path), width=Inches(6.5))
        p_cap = doc_target.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(f"Hình — {idx}.2: Sơ đồ hoạt động {uc_id} — {title}")
        fmt(r_cap, size=8.5, italic=True, color=(100, 116, 139))
    else:
        print(f"Warning: Activity image {act_img_path} not found!")

    # Steps Table
    p_tbl_act_lbl = doc_target.add_paragraph()
    p_tbl_act_lbl.paragraph_format.space_after = Pt(4)
    r = p_tbl_act_lbl.add_run("Bảng chi tiết các bước hoạt động:")
    fmt(r, size=9.5, bold=True, color=(30, 41, 59))
    
    tbl_desc = doc_target.add_table(rows=1, cols=4)
    set_table_col_widths(tbl_desc, [0.45, 1.4, 1.8, 2.85])
    headers = ['#', 'Phân làn (Partition)', 'Loại nút / Quyết định', 'Nội dung hành vi / Điều kiện']
    for c, h in zip(tbl_desc.rows[0].cells, headers):
        set_cell(c, h, bold=True, size=8.5, color=(255, 255, 255))
        shade_cell(c, '334155')
        
    step_counter = 1
    for st_part, st_type, st_text, st_cond in act_case['steps_def']:
        if st_type in ['action', 'decision', 'else']:
            row = tbl_desc.add_row()
            set_cell(row.cells[0], str(step_counter), bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell(row.cells[1], st_part, bold=True, size=8)
            
            type_label = 'Hành động' if st_type == 'action' else ('Nút quyết định' if st_type == 'decision' else 'Nhánh điều kiện')
            set_cell(row.cells[2], type_label, bold=False, size=8)
            
            if st_type == 'else':
                shade_cell(row.cells[2], 'FEF2F2')
                shade_cell(row.cells[3], 'FEF2F2')
                set_cell(row.cells[3], f"[{st_cond}] -> {st_text}", bold=False, size=8, color=(185, 28, 28))
            else:
                set_cell(row.cells[3], st_text, bold=False, size=8)
                
            step_counter += 1
            
    p_sp2 = doc_target.add_paragraph()
    p_sp2.paragraph_format.space_after = Pt(6)
    
    # PlantUML Activity Code
    h3_act_code = doc_target.add_heading(level=2)
    r = h3_act_code.add_run(f"Mã PlantUML sơ đồ hoạt động của {title}")
    fmt(r, size=11, bold=True, color=(30, 41, 59))
    
    puml_act_str = generate_activity_plantuml(act_case['title'], idx, act_case['partitions'], act_case['steps_def'])
    add_code_block(doc_target, puml_act_str)

# Section Final: Audit Notes
doc_target.add_page_break()
h1_end = doc_target.add_heading(level=1)
r = h1_end.add_run('Ghi chú kiểm chứng & Nguyên tắc xử lý ngoại lệ')
fmt(r, size=14, bold=True, color=(15, 23, 42))

p_end = doc_target.add_paragraph()
p_end.paragraph_format.line_spacing = 1.15
r = p_end.add_run(
    '1. Sơ đồ tuần tự và sơ đồ hoạt động tuân thủ nguyên tắc không tự sinh nghiệp vụ mới ngoài phạm vi được quy định trong tài liệu "Đặc tả UC_done.docx".\n'
    '2. Nhánh ngoại lệ (E1, E2) tập trung biểu diễn việc gián đoạn luồng chính, thông báo lỗi hệ thống/CSDL/API và hủy bỏ phiên thao tác để bảo toàn tính nhất quán dữ liệu.\n'
    '3. Nhánh thay thế (A1, A2, A3) biểu diễn các phản hồi điều hướng của hệ thống khi dữ liệu đầu vào không hợp lệ hoặc người dùng không đủ quyền thao tác.\n'
    '4. Đối với các UC tích hợp AI (UC-010, UC-011, UC-012), các xử lý fallback khi AI timeout, rate limit hoặc response sai định dạng được biểu diễn rõ ràng nhằm đảm bảo hệ thống không bịa số liệu và không tự động áp đặt quyết định vận hành.'
)
fmt(r, size=10, color=(30, 41, 59))

output_path = Path('Sơ_đồ_tuần_tự_hoạt_động.docx')
try:
    doc_target.save(output_path)
    print('====================================================')
    print(f'SUCCESS! File saved to: {output_path.resolve()}')
    print('====================================================')
except PermissionError:
    alt_path = Path('Sơ_đồ_tuần_tự_hoạt_động_fixed.docx')
    doc_target.save(alt_path)
    print('====================================================')
    print(f'Sơ_đồ_tuần_tự_hoạt_động.docx đang mở trong Word! Đã lưu tạm sang: {alt_path.resolve()}')
    print('====================================================')
