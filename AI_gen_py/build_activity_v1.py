import sys
import os
import textwrap
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

ASSETS_DIR = Path('act_assets')
ASSETS_DIR.mkdir(exist_ok=True)

# ---------------------------------------------------------
# Helper functions for python-docx styling & formatting
# ---------------------------------------------------------
def fmt(run, size=10, bold=None, italic=None, color=None, font_name='Calibri'):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rPr.append(rFonts)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_cell(cell, text, bold=False, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    fmt(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def set_table_col_widths(table, widths_in_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in_inches):
            cell.width = Inches(w)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    set_table_col_widths(tbl, [6.5])
    c = tbl.cell(0, 0)
    shade_cell(c, 'F8FAFC')
    p = c.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(code_str)
    fmt(r, size=7.5, bold=False, color=(30, 41, 59), font_name='Consolas')
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

# ---------------------------------------------------------
# Image Generator Engine for Activity Diagrams using Pillow
# ---------------------------------------------------------
def draw_activity_diagram(title, partitions, nodes, output_filename):
    width = 1600
    n_part = len(partitions)
    col_width = width / n_part
    
    header_h = 100
    row_h = 95
    footer_h = 60
    
    # Calculate image height based on nodes count
    total_h = max(750, header_h + len(nodes) * row_h + footer_h)
    
    img = Image.new('RGB', (width, total_h), 'white')
    d = ImageDraw.Draw(img)
    
    try:
        font_title = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 24)
        font_part = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 22)
        font_bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 19)
        font_normal = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 17)
        font_small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 15)
    except OSError:
        font_title = font_part = font_bold = font_normal = font_small = ImageFont.load_default()
        
    # Title box
    d.rounded_rectangle((20, 15, width - 20, 65), radius=8, fill='#F1F5F9', outline='#0F172A', width=2)
    d.text((35, 25), f'Sơ đồ hoạt động UML 2.5: {title}', font=font_title, fill='#0F172A')
    
    # Draw Swimlanes (Partitions)
    part_xs = []
    for i, p_name in enumerate(partitions):
        x1 = i * col_width
        x2 = (i + 1) * col_width
        mid_x = (x1 + x2) / 2
        part_xs.append(mid_x)
        
        # Header background
        d.rectangle((x1, 75, x2, 120), fill='#E2E8F0', outline='#64748B', width=2)
        bbox = d.textbbox((0, 0), f'|{p_name}|', font=font_part)
        tw = bbox[2] - bbox[0]
        d.text((mid_x - tw/2, 88), f'|{p_name}|', font=font_part, fill='#0F172A')
        
        # Partition divider vertical line
        if i > 0:
            for ly in range(120, total_h - 30, 12):
                d.line((x1, ly, x1, min(ly + 6, total_h - 30)), fill='#CBD5E1', width=2)
                
    # Draw Nodes & Connecting Arrows
    y_curr = 155
    node_coords = []
    
    for i, item in enumerate(nodes):
        node_type = item[0]  # 'start', 'action', 'decision', 'stop'
        part_idx = item[1]
        text = item[2] if len(item) > 2 else ''
        cond_text = item[3] if len(item) > 3 else ''
        
        x = part_xs[part_idx]
        y = y_curr
        node_coords.append((node_type, x, y, text))
        
        if node_type == 'start':
            # Solid black circle
            d.ellipse((x - 18, y - 18, x + 18, y + 18), fill='#0F172A', outline='#0F172A')
            d.text((x + 28, y - 10), 'Bắt đầu', font=font_bold, fill='#0F172A')
            y_curr += 85
            
        elif node_type == 'stop':
            # Double circle (Bullseye)
            d.ellipse((x - 20, y - 20, x + 20, y + 20), fill='white', outline='#0F172A', width=3)
            d.ellipse((x - 12, y - 12, x + 12, y + 12), fill='#0F172A', outline='#0F172A')
            d.text((x + 28, y - 10), 'Kết thúc', font=font_bold, fill='#0F172A')
            y_curr += 85
            
        elif node_type == 'action':
            # Action Rounded Rectangle
            lines = textwrap.wrap(text, width=28)
            lbl = '\n'.join(lines)
            bbox = d.multiline_textbbox((0, 0), lbl, font=font_normal, spacing=3)
            tw = max(180, (bbox[2] - bbox[0]) + 36)
            th = max(46, (bbox[3] - bbox[1]) + 20)
            
            d.rounded_rectangle((x - tw/2, y - th/2, x + tw/2, y + th/2), radius=12, fill='#EFF6FF', outline='#2563EB', width=2)
            d.multiline_text((x - (bbox[2]-bbox[0])/2, y - (bbox[3]-bbox[1])/2), lbl, font=font_normal, fill='#1E3A8A', align='center', spacing=3)
            y_curr += th + 45
            
        elif node_type == 'decision':
            # Diamond Node
            lines = textwrap.wrap(text, width=24)
            lbl = '\n'.join(lines)
            bbox = d.multiline_textbbox((0, 0), lbl, font=font_small, spacing=2)
            tw = max(170, (bbox[2] - bbox[0]) + 40)
            th = max(55, (bbox[3] - bbox[1]) + 24)
            
            points = [(x, y - th/2 - 5), (x + tw/2 + 10, y), (x, y + th/2 + 5), (x - tw/2 - 10, y)]
            d.polygon(points, fill='#FEF3C7', outline='#D97706', width=2)
            d.multiline_text((x - (bbox[2]-bbox[0])/2, y - (bbox[3]-bbox[1])/2), lbl, font=font_small, fill='#78350F', align='center', spacing=2)
            y_curr += th + 55

    # Draw Connecting Arrows between sequential nodes
    for i in range(len(node_coords) - 1):
        t1, x1, y1, _ = node_coords[i]
        t2, x2, y2, _ = node_coords[i + 1]
        
        # Calculate offset from node border
        y_start = y1 + 22
        y_end = y2 - 25
        
        if x1 == x2:
            d.line((x1, y_start, x2, y_end), fill='#334155', width=2)
            d.polygon([(x2, y_end), (x2 - 6, y_end - 10), (x2 + 6, y_end - 10)], fill='#334155')
        else:
            # L-shaped connector for swimlane transition
            mid_y = (y_start + y_end) / 2
            d.line((x1, y_start, x1, mid_y), fill='#334155', width=2)
            d.line((x1, mid_y, x2, mid_y), fill='#334155', width=2)
            d.line((x2, mid_y, x2, y_end), fill='#334155', width=2)
            d.polygon([(x2, y_end), (x2 - 6, y_end - 10), (x2 + 6, y_end - 10)], fill='#334155')
            
        # Add condition label if decision
        if t1 == 'decision':
            cond = nodes[i][3] if len(nodes[i]) > 3 else ''
            if cond:
                d.rectangle((x1 + 10, y_start + 2, x1 + 140, y_start + 22), fill='white')
                d.text((x1 + 14, y_start + 3), f'[{cond}]', font=font_small, fill='#B45309')
                
    out_path = ASSETS_DIR / output_filename
    img.save(out_path)
    return str(out_path)

# Helper to generate PlantUML Activity Code String
def generate_activity_plantuml(title, idx_num, partitions, steps_def):
    out = [
        '@startuml',
        f'title {idx_num}. Sơ đồ hoạt động của {title}',
        'skinparam ActivityFontSize 13',
        'skinparam ActivityBorderColor #1E293B',
        'skinparam ActivityBackgroundColor #F8FAFC',
        'skinparam ActivityDiamondBorderColor #D97706',
        'skinparam ActivityDiamondBackgroundColor #FEF3C7'
    ]
    
    curr_part = None
    for part, action_type, text, cond in steps_def:
        if part != curr_part:
            out.append(f'|{part}|')
            curr_part = part
            
        if action_type == 'start':
            out.append('start')
        elif action_type == 'stop':
            out.append('stop')
        elif action_type == 'action':
            out.append(f':{text};')
        elif action_type == 'decision':
            out.append(f'if ({text}?) then ([Đúng])')
        elif action_type == 'else':
            out.append(f'else ([{cond}])')
            out.append(f'  :{text};')
        elif action_type == 'endif':
            out.append('endif')
            
    out.append('@enduml')
    return '\n'.join(out)

# ---------------------------------------------------------
# Define complete dataset for all 12 Use Cases
# ---------------------------------------------------------
ACTIVITY_CASES = [
    {
        'idx': 1,
        'title': 'Đăng nhập và phân quyền',
        'uc_id': 'UC-001',
        'overview': 'Mô tả tiến trình xác thực thông tin tài khoản người dùng, kiểm tra phân quyền truy cập chức năng và xử lý các nhánh từ chối đăng nhập (A1), từ chối vượt quyền (A2) và lỗi kết nối CSDL (E1).',
        'partitions': ['Người dùng', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Nhập tên đăng nhập & mật khẩu'),
            ('action', 1, 'Tiếp nhận & đối chiếu CSDL'),
            ('decision', 1, 'Thông tin hợp lệ & CSDL sẵn sàng?', 'Hợp lệ'),
            ('action', 1, 'Xác định vai trò & cấp phiên làm việc'),
            ('decision', 1, 'Có quyền truy cập chức năng?', 'Đủ quyền'),
            ('action', 1, 'Cấp menu chức năng phù hợp'),
            ('action', 0, 'Truy cập màn hình theo vai trò'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Người dùng', 'start', '', ''),
            ('Người dùng', 'action', 'Nhập tên đăng nhập và mật khẩu', ''),
            ('Hệ thống', 'action', 'Đối chiếu thông tin tài khoản với CSDL', ''),
            ('Hệ thống', 'decision', 'Thông tin đăng nhập hợp lệ', ''),
            ('Hệ thống', 'action', 'Xác định vai trò và cấp phiên làm việc', ''),
            ('Hệ thống', 'decision', 'Có quyền truy cập chức năng yêu cầu', ''),
            ('Hệ thống', 'action', 'Cấp menu chức năng phù hợp theo vai trò', ''),
            ('Người dùng', 'action', 'Đăng nhập thành công và truy cập hệ thống', ''),
            ('Người dùng', 'stop', '', ''),
            ('Hệ thống', 'else', 'Báo lỗi không thể xác thực (A1) hoặc lỗi CSDL (E1)', 'Không hợp lệ / Lỗi CSDL'),
            ('Người dùng', 'action', 'Xem thông báo lỗi và kết thúc Use Case', ''),
            ('Người dùng', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 2,
        'title': 'Quản lý tài khoản và phân quyền',
        'uc_id': 'UC-002',
        'overview': 'Tiến trình quản lý thông tin tài khoản và cấp quyền theo vai trò do Quản lý thực hiện. Bao gồm kiểm tra quyền Quản lý (A2), kiểm tra dữ liệu hợp lệ (A1) và xử lý lỗi lưu CSDL (E1).',
        'partitions': ['Quản lý', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Chọn chức năng & nhập dữ liệu tài khoản/quyền'),
            ('action', 1, 'Kiểm tra quyền Quản lý thao tác'),
            ('decision', 1, 'Actor có đủ quyền quản trị?', 'Đủ quyền'),
            ('action', 1, 'Kiểm tra định dạng dữ liệu tài khoản'),
            ('decision', 1, 'Dữ liệu hợp lệ & không trùng?', 'Hợp lệ'),
            ('action', 1, 'Lưu cập nhật tài khoản/quyền vào CSDL'),
            ('action', 0, 'Nhận thông báo cập nhật thành công'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Chọn thao tác & cung cấp thay đổi tài khoản/quyền', ''),
            ('Hệ thống', 'action', 'Kiểm tra quyền Quản lý & kiểm tra định dạng dữ liệu', ''),
            ('Hệ thống', 'decision', 'Quản lý đủ quyền & Dữ liệu hợp lệ', ''),
            ('Hệ thống', 'action', 'Lưu thay đổi tài khoản/quyền vào CSDL', ''),
            ('Quản lý', 'action', 'Nhận thông báo cập nhật thành công', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Từ chối thao tác (A2) hoặc Báo lỗi dữ liệu không hợp lệ (A1)', 'Không hợp lệ'),
            ('Quản lý', 'action', 'Xem thông báo từ chối / Yêu cầu điều chỉnh lại', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 3,
        'title': 'Quản lý khu vực, vị trí đỗ và loại xe',
        'uc_id': 'UC-003',
        'overview': 'Tiến trình cấu hình danh mục khu vực bãi xe, số lượng vị trí đỗ và các loại phương tiện. Bao gồm rẽ nhánh kiểm tra dữ liệu trùng/thiếu (A1) và xử lý lỗi kết nối CSDL (E1).',
        'partitions': ['Quản lý', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Chọn danh mục khu vực/vị trí/loại xe & nhập dữ liệu'),
            ('action', 1, 'Tiếp nhận & kiểm tra dữ liệu danh mục'),
            ('decision', 1, 'Dữ liệu hợp lệ & không bị trùng?', 'Hợp lệ'),
            ('action', 1, 'Lưu dữ liệu danh mục mới vào CSDL'),
            ('action', 0, 'Xem danh mục mới & nhận thông báo thành công'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Nhập thông tin khu vực, vị trí đỗ hoặc loại xe mới', ''),
            ('Hệ thống', 'action', 'Kiểm tra tính hợp lệ & trùng lặp danh mục', ''),
            ('Hệ thống', 'decision', 'Dữ liệu danh mục hợp lệ', ''),
            ('Hệ thống', 'action', 'Cập nhật CSDL và hiển thị danh mục mới', ''),
            ('Quản lý', 'action', 'Xem thông báo thành công và danh mục đã cập nhật', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Báo lỗi dữ liệu trùng/thiếu (A1) hoặc Lỗi lưu CSDL (E1)', 'Lỗi dữ liệu / CSDL'),
            ('Quản lý', 'action', 'Xem thông báo lỗi & điều chỉnh thông tin', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 4,
        'title': 'Ghi nhận xe vào',
        'uc_id': 'UC-004',
        'overview': 'Tiến trình xử lý xe vào bãi: quét thẻ, kiểm tra vị trí đỗ trống, ghi nhận lượt vào và mở barie. Xử lý các luồng bãi xe hết chỗ (A1), thẻ bị khóa/dùng trùng (A2) và lỗi CSDL (E1).',
        'partitions': ['Nhân viên bãi xe', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Quét thẻ vé / Nhập thông tin xe vào'),
            ('action', 1, 'Kiểm tra trạng thái thẻ vé & bãi đỗ'),
            ('decision', 1, 'Thẻ hợp lệ & Bãi đỗ còn chỗ trống?', 'Còn chỗ & Thẻ hợp lệ'),
            ('action', 1, 'Ghi nhận lượt xe vào & gán vị trí đỗ vào CSDL'),
            ('action', 0, 'In vé / Mở barie cho xe vào bãi'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Nhân viên bãi xe', 'start', '', ''),
            ('Nhân viên bãi xe', 'action', 'Quét thẻ vé và nhập thông tin phương tiện vào', ''),
            ('Hệ thống', 'action', 'Kiểm tra trạng thái thẻ vé và vị trí đỗ còn trống', ''),
            ('Hệ thống', 'decision', 'Thẻ hợp lệ và Bãi đỗ còn vị trí trống', ''),
            ('Hệ thống', 'action', 'Ghi nhận lượt vào CSDL và cập nhật trạng thái chỗ đỗ', ''),
            ('Nhân viên bãi xe', 'action', 'Xác nhận lượt vào, in vé và mở barie cho xe', ''),
            ('Nhân viên bãi xe', 'stop', '', ''),
            ('Hệ thống', 'else', 'Báo hết chỗ đỗ (A1) hoặc Thẻ không hợp lệ (A2)', 'Hết chỗ / Thẻ lỗi'),
            ('Nhân viên bãi xe', 'action', 'Thông báo cho chủ xe & từ chối lượt vào', ''),
            ('Nhân viên bãi xe', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 5,
        'title': 'Ghi nhận xe ra và tính phí',
        'uc_id': 'UC-005',
        'overview': 'Tiến trình hoàn tất lượt gửi xe: quẹt thẻ ra, truy xuất lượt vào, tính phí theo loại xe/thời gian, thu tiền và giải phóng chỗ. Xử lý không thấy lượt vào (A1), lỗi thời gian (A2), lỗi bảng giá (A3) và lỗi CSDL (E1).',
        'partitions': ['Nhân viên bãi xe', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Quét thẻ vé / Nhập thông tin xe ra'),
            ('action', 1, 'Tìm lượt xe vào tương ứng trong CSDL'),
            ('decision', 1, 'Tìm thấy lượt vào & Thời gian hợp lệ?', 'Tìm thấy & Hợp lệ'),
            ('action', 1, 'Tính phí gửi xe theo loại xe và thời gian'),
            ('action', 0, 'Thu tiền gửi xe & Xác nhận thanh toán'),
            ('action', 1, 'Ghi nhận xe ra & Giải phóng vị trí đỗ CSDL'),
            ('action', 0, 'Mở barie cho xe rời bãi'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Nhân viên bãi xe', 'start', '', ''),
            ('Nhân viên bãi xe', 'action', 'Quét thẻ vé xe ra để hệ thống tra cứu lượt vào', ''),
            ('Hệ thống', 'action', 'Tìm lượt xe vào và gọi Bảng giá tính phí gửi xe', ''),
            ('Hệ thống', 'decision', 'Tìm thấy lượt vào và Tính được mức phí', ''),
            ('Nhân viên bãi xe', 'action', 'Thu số tiền phí hiển thị và bấm xác nhận thanh toán', ''),
            ('Hệ thống', 'action', 'Ghi nhận hoàn tất xe ra & giải phóng vị trí đỗ trong CSDL', ''),
            ('Nhân viên bãi xe', 'action', 'Mở barie cho phương tiện xuất bãi', ''),
            ('Nhân viên bãi xe', 'stop', '', ''),
            ('Hệ thống', 'else', 'Không tìm thấy lượt vào (A1) hoặc Lỗi tính phí (A3)', 'Không tìm thấy / Lỗi giá'),
            ('Nhân viên bãi xe', 'action', 'Thông báo không thể hoàn tất & xử lý sự cố thủ công', ''),
            ('Nhân viên bãi xe', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 6,
        'title': 'Theo dõi chỗ trống theo khu vực',
        'uc_id': 'UC-006',
        'overview': 'Tiến trình truy xuất và tổng hợp bản đồ chỗ trống bãi đỗ theo khu vực. Xử lý trường hợp không có dữ liệu khu vực (A1), vị trí chưa xác định (A2) và lỗi kết nối CSDL (E1).',
        'partitions': ['Người dùng', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Chọn xem tình trạng chỗ trống theo khu vực'),
            ('action', 1, 'Truy xuất trạng thái vị trí đỗ từ CSDL'),
            ('decision', 1, 'Có dữ liệu vị trí các khu vực?', 'Có dữ liệu'),
            ('action', 1, 'Tổng hợp số chỗ trống & lập sơ đồ hiển thị'),
            ('action', 0, 'Xem sơ đồ & số lượng chỗ trống các khu vực'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Người dùng', 'start', '', ''),
            ('Người dùng', 'action', 'Chọn chức năng xem tình trạng chỗ đỗ theo khu vực', ''),
            ('Hệ thống', 'action', 'Truy xuất CSDL & tổng hợp số chỗ trống từng khu vực', ''),
            ('Hệ thống', 'decision', 'Có dữ liệu khu vực phù hợp', ''),
            ('Hệ thống', 'action', 'Hiển thị sơ đồ chỗ trống trực quan cho người dùng', ''),
            ('Người dùng', 'action', 'Theo dõi bản đồ vị trí đỗ trống', ''),
            ('Người dùng', 'stop', '', ''),
            ('Hệ thống', 'else', 'Không có dữ liệu khu vực (A1) hoặc Lỗi CSDL (E1)', 'Không có dữ liệu / Lỗi CSDL'),
            ('Người dùng', 'action', 'Xem thông báo dữ liệu không sẵn có', ''),
            ('Người dùng', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 7,
        'title': 'Tra cứu lượt gửi xe',
        'uc_id': 'UC-007',
        'overview': 'Tiến trình tìm kiếm lịch sử lượt xe gửi theo biển số hoặc khoảng thời gian. Xử lý các nhánh tiêu chí nhập sai định dạng (A1), không có kết quả phù hợp (A2) và lỗi CSDL (E1).',
        'partitions': ['Người dùng', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Nhập biển số xe hoặc chọn khoảng thời gian tra cứu'),
            ('action', 1, 'Kiểm tra tiêu chí & truy xuất CSDL'),
            ('decision', 1, 'Tiêu chí hợp lệ & Tìm thấy kết quả?', 'Tìm thấy'),
            ('action', 1, 'Tổng hợp danh sách lượt gửi phù hợp'),
            ('action', 0, 'Xem bảng kết quả tra cứu lượt gửi xe'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Người dùng', 'start', '', ''),
            ('Người dùng', 'action', 'Nhập tiêu chí tra cứu (Biển số xe và/hoặc Thời gian)', ''),
            ('Hệ thống', 'action', 'Kiểm tra tiêu chí & truy xuất lượt xe phù hợp từ CSDL', ''),
            ('Hệ thống', 'decision', 'Tiêu chí hợp lệ & Tìm thấy kết quả', ''),
            ('Hệ thống', 'action', 'Trả về danh sách chi tiết các lượt gửi xe', ''),
            ('Người dùng', 'action', 'Theo dõi kết quả tra cứu hiển thị trên màn hình', ''),
            ('Người dùng', 'stop', '', ''),
            ('Hệ thống', 'else', 'Tiêu chí sai (A1) hoặc Không có kết quả phù hợp (A2)', 'Không tìm thấy'),
            ('Người dùng', 'action', 'Xem thông báo không tìm thấy kết quả', ''),
            ('Người dùng', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 8,
        'title': 'Quản lý vé tháng hoặc khách quen',
        'uc_id': 'UC-008',
        'overview': 'Tiến trình đăng ký mới, điều chỉnh thông tin hoặc gia hạn vé tháng/khách quen. Bao gồm kiểm tra thông tin hợp lệ (A1), vé không tồn tại/hết hạn (A2) và lỗi lưu CSDL (E1).',
        'partitions': ['Quản lý', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Nhập/sửa thông tin vé tháng hoặc khách quen'),
            ('action', 1, 'Kiểm tra định dạng & đối chiếu CSDL'),
            ('decision', 1, 'Thông tin hợp lệ & Vé tồn tại?', 'Hợp lệ'),
            ('action', 1, 'Lưu thông tin vé vào CSDL & Cập nhật trạng thái'),
            ('action', 0, 'Nhận thông báo cập nhật vé thành công'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Cung cấp thông tin đăng ký mới hoặc gia hạn vé tháng', ''),
            ('Hệ thống', 'action', 'Kiểm tra quy tắc nghiệp vụ vé & tồn tại trong CSDL', ''),
            ('Hệ thống', 'decision', 'Thông tin vé hợp lệ & Đúng quy tắc', ''),
            ('Hệ thống', 'action', 'Lưu dữ liệu vé vào CSDL và gửi xác nhận', ''),
            ('Quản lý', 'action', 'Xem thông báo cập nhật thông tin vé thành công', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Dữ liệu vé sai (A1) hoặc Vé không tồn tại/hết hạn (A2)', 'Lỗi dữ liệu / Không tồn tại'),
            ('Quản lý', 'action', 'Xem thông báo lỗi & kiểm tra lại thông tin vé', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 9,
        'title': 'Xem thống kê vận hành',
        'uc_id': 'UC-009',
        'overview': 'Tiến trình tổng hợp lưu lượng lượt xe, doanh thu và xác định khung giờ cao điểm theo khoảng thời gian. Xử lý dữ liệu rỗng (A1), mốc thời gian sai (A2) và lỗi CSDL (E1).',
        'partitions': ['Quản lý', 'Hệ thống'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Chọn khoảng thời gian & bộ lọc xem thống kê'),
            ('action', 1, 'Kiểm tra mốc thời gian & truy xuất CSDL'),
            ('decision', 1, 'Thời gian hợp lệ & Có đủ dữ liệu?', 'Đủ dữ liệu'),
            ('action', 1, 'Tổng hợp lưu lượng, doanh thu & khung giờ cao điểm'),
            ('action', 0, 'Xem báo cáo biểu đồ thống kê vận hành'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Chọn khoảng thời gian (ngày/tuần/tháng) cần xem thống kê', ''),
            ('Hệ thống', 'action', 'Truy xuất CSDL lượt gửi xe & tính toán lưu lượng, doanh thu', ''),
            ('Hệ thống', 'decision', 'Thời gian hợp lệ & Đủ dữ liệu thống kê', ''),
            ('Hệ thống', 'action', 'Hiển thị báo cáo biểu đồ thống kê vận hành chi tiết', ''),
            ('Quản lý', 'action', 'Theo dõi báo cáo lưu lượng và doanh thu', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Dữ liệu rỗng (A1) hoặc Thời gian không hợp lệ (A2)', 'Dữ liệu rỗng / Sai mốc'),
            ('Quản lý', 'action', 'Xem thông báo không đủ dữ liệu thống kê', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 10,
        'title': 'AI sinh báo cáo lưu lượng',
        'uc_id': 'UC-010',
        'overview': 'Tiến trình tự động tạo văn bản nhận xét báo cáo lưu lượng ngày/tuần bằng AI. Bao gồm kiểm tra phạm vi (A2), kiểm tra dữ liệu thống kê (A1), xử lý AI timeout/rate limit (E1) và response sai format (E2).',
        'partitions': ['Quản lý', 'Hệ thống', 'Nhà cung cấp AI'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Yêu cầu tạo báo cáo AI (ngày hoặc tuần)'),
            ('action', 1, 'Lấy dữ liệu thống kê & Xây dựng prompt guardrails'),
            ('decision', 1, 'Phạm vi hợp lệ & Đủ dữ liệu thống kê?', 'Đủ dữ liệu'),
            ('action', 2, 'Gửi prompt gọi API mô hình AI (LLM)'),
            ('action', 2, 'Xử lý suy luận & trả về kết quả báo cáo'),
            ('action', 1, 'Kiểm tra format & Đính kèm cảnh báo kiểm chứng'),
            ('action', 0, 'Xem báo cáo AI & tự kiểm chứng số liệu'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Yêu cầu tạo báo cáo phân tích AI theo ngày hoặc tuần', ''),
            ('Hệ thống', 'action', 'Lấy dữ liệu thống kê & gửi prompt kèm guardrails tới AI', ''),
            ('Nhà cung cấp AI', 'action', 'Xử lý mô hình ngôn ngữ & trả về văn bản nhận xét', ''),
            ('Hệ thống', 'decision', 'AI phản hồi thành công & Đúng định dạng', ''),
            ('Hệ thống', 'action', 'Hiển thị báo cáo AI kèm cảnh báo Quản lý cần kiểm chứng', ''),
            ('Quản lý', 'action', 'Xem nội dung báo cáo và tự kiểm chứng số liệu', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Thiếu dữ liệu (A1) hoặc AI timeout/lỗi format (E1, E2)', 'AI Lỗi / Thiếu dữ liệu'),
            ('Quản lý', 'action', 'Xem thông báo AI không khả dụng & không tự bịa số liệu', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 11,
        'title': 'AI hỏi đáp dữ liệu bãi xe',
        'uc_id': 'UC-011',
        'overview': 'Tiến trình giải đáp câu hỏi quản trị bằng AI dựa trên ngữ cảnh dữ liệu bãi xe. Bao gồm kiểm tra vượt phân quyền (A2), không có dữ liệu phù hợp (A1) và xử lý AI lỗi/không khả dụng (E1).',
        'partitions': ['Quản lý', 'Hệ thống', 'Nhà cung cấp AI'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Nhập câu hỏi phân tích dữ liệu bãi xe'),
            ('action', 1, 'Kiểm tra quyền & Lấy ngữ cảnh dữ liệu CSDL'),
            ('decision', 1, 'Đủ quyền & Có dữ liệu ngữ cảnh phù hợp?', 'Đủ ngữ cảnh'),
            ('action', 2, 'Gửi context + prompt câu hỏi tới API AI'),
            ('action', 2, 'Sinh câu trả lời phân tích từ ngữ cảnh'),
            ('action', 1, 'Kiểm tra kết quả & Đính kèm khuyến cáo kiểm chứng'),
            ('action', 0, 'Xem câu trả lời AI & tự xác minh'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Nhập câu hỏi quản trị cần AI giải đáp phân tích', ''),
            ('Hệ thống', 'action', 'Kiểm tra quyền Quản lý & lấy dữ liệu ngữ cảnh gửi AI', ''),
            ('Nhà cung cấp AI', 'action', 'Phân tích ngữ cảnh & tạo câu trả lời giải đáp', ''),
            ('Hệ thống', 'decision', 'Có dữ liệu ngữ cảnh & AI trả lời thành công', ''),
            ('Hệ thống', 'action', 'Hiển thị câu trả lời AI kèm lưu ý kiểm chứng', ''),
            ('Quản lý', 'action', 'Theo dõi câu trả lời phân tích từ AI', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Vượt quyền (A2) hoặc AI không khả dụng/lỗi (E1)', 'Vượt quyền / AI lỗi'),
            ('Quản lý', 'action', 'Xem thông báo từ chối / Lỗi dịch vụ AI', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    },
    {
        'idx': 12,
        'title': 'AI gợi ý bố trí nhân sự',
        'uc_id': 'UC-012',
        'overview': 'Tiến trình tham khảo gợi ý phân bổ ca làm việc nhân sự theo khung giờ cao điểm từ AI. Bao gồm kiểm tra thiếu dữ liệu cao điểm (A1), gợi ý rỗng (A2) và AI timeout/rate limit (E1).',
        'partitions': ['Quản lý', 'Hệ thống', 'Nhà cung cấp AI'],
        'nodes': [
            ('start', 0),
            ('action', 0, 'Yêu cầu gợi ý phương án bố trí nhân sự'),
            ('action', 1, 'Truy xuất dữ liệu khung giờ cao điểm từ CSDL'),
            ('decision', 1, 'Đủ dữ liệu cao điểm lịch sử?', 'Đủ dữ liệu'),
            ('action', 2, 'Gửi dữ liệu cao điểm tới AI xin gợi ý'),
            ('action', 2, 'Sinh phương án gợi ý phân bổ ca nhân sự'),
            ('action', 1, 'Kiểm tra & Đính kèm lưu ý chỉ mang tính tham khảo'),
            ('action', 0, 'Xem phương án gợi ý AI & Tự ra quyết định'),
            ('stop', 0)
        ],
        'steps_def': [
            ('Quản lý', 'start', '', ''),
            ('Quản lý', 'action', 'Yêu cầu AI đưa ra phương án gợi ý bố trí nhân sự ca làm', ''),
            ('Hệ thống', 'action', 'Lấy dữ liệu khung giờ cao điểm lịch sử gửi cho AI', ''),
            ('Nhà cung cấp AI', 'action', 'Phân tích lưu lượng & đề xuất gợi ý phân bổ ca', ''),
            ('Hệ thống', 'decision', 'Đủ dữ liệu cao điểm & AI đề xuất thành công', ''),
            ('Hệ thống', 'action', 'Hiển thị gợi ý nhân sự (chỉ mang tính tham khảo)', ''),
            ('Quản lý', 'action', 'Xem gợi ý và tự quyết định phương án vận hành', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'else', 'Thiếu dữ liệu cao điểm (A1) hoặc AI lỗi/timeout (E1)', 'Thiếu dữ liệu / AI lỗi'),
            ('Quản lý', 'action', 'Xem thông báo lỗi AI; không tự động thay đổi vận hành', ''),
            ('Quản lý', 'stop', '', ''),
            ('Hệ thống', 'endif', '', '')
        ]
    }
]

# ---------------------------------------------------------
# Build Act_Diagram_1.0.docx
# ---------------------------------------------------------
print('Starting document assembly for Act_Diagram_1.0.docx...')
doc_target = Document()

# Set standard margins (1 inch = 72 pt)
for s in doc_target.sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

# Document Title
p_title = doc_target.add_paragraph()
r_title = p_title.add_run('HỆ THỐNG QUẢN LÝ BÃI XE THÔNG MINH')
fmt(r_title, size=18, bold=True, color=(15, 23, 42))

p_sub = doc_target.add_paragraph()
r_sub = p_sub.add_run('TÀI LIỆU SƠ ĐỒ HOẠT ĐỘNG (ACTIVITY DIAGRAMS v1.0)')
fmt(r_sub, size=14, bold=True, color=(37, 99, 235))

p_sub2 = doc_target.add_paragraph()
r_sub2 = p_sub2.add_run('Biểu diễn luồng tiến trình Swimlanes (Partition) & Nhánh quyết định cho 12 Use Case tiêu chuẩn')
fmt(r_sub2, size=10.5, italic=True, color=(71, 85, 105))

# Document Metadata Table
tbl_meta = doc_target.add_table(rows=4, cols=2)
set_table_col_widths(tbl_meta, [2.0, 4.5])
meta_data = [
    ('Nguồn đặc tả', 'Đặc tả UC_done.docx (Nhóm 10 — Đinh Tiến Mạnh, Tạ Văn Đức)'),
    ('Phạm vi hệ thống', '12 Use Case tiêu chuẩn (UC-001 đến UC-012)'),
    ('Quy ước sơ đồ', 'Activity Diagram chuẩn UML 2.5 với Swimlanes, Action nodes & Decision nodes'),
    ('Phiên bản tài liệu', 'v1.0 — Chuẩn hóa tiêu đề theo mẫu "1. Sơ đồ hoạt động của <Tên UC>"')
]
for i, (k, v) in enumerate(meta_data):
    row = tbl_meta.rows[i]
    set_cell(row.cells[0], k, bold=True, size=9, color=(15, 23, 42))
    shade_cell(row.cells[0], 'F1F5F9')
    set_cell(row.cells[1], v, bold=False, size=9)

p_space = doc_target.add_paragraph()
p_space.paragraph_format.space_after = Pt(12)

# Section 1: Conventions
h1_1 = doc_target.add_heading(level=1)
r = h1_1.add_run('Hướng dẫn & Quy ước Sơ đồ hoạt động (Activity Diagram)')
fmt(r, size=14, bold=True, color=(15, 23, 42))

p_conv = doc_target.add_paragraph()
p_conv.paragraph_format.line_spacing = 1.15
r = p_conv.add_run(
    'Sơ đồ hoạt động (Activity Diagram) thể hiện dòng điều khiển (Control Flow) và dòng dữ liệu giữa Tác nhân (Actor) và các thành phần Hệ thống Backend/CSDL/Dịch vụ AI:\n'
    '• Swimlanes (Partition): Phân làn trách nhiệm giữa Người dùng/Quản lý/Nhân viên và Hệ thống xử lý.\n'
    '• Nút hình tròn đen (●): Điểm bắt đầu (Start Node).\n'
    '• Nút hai vòng tròn (◉): Điểm kết thúc (Stop/End Node).\n'
    '• Nút hình chữ nhật bo góc: Hành vi/Thao tác xử lý (Action State).\n'
    '• Nút hình thoi (◇): Nút quyết định (Decision Node) rẽ nhánh sang Luồng chính, Luồng thay thế (A1, A2) hoặc Luồng ngoại lệ (E1, E2).\n'
    '• Đi kèm mỗi sơ đồ là Mã PlantUML Activity Diagram Beta syntax sẵn sàng nạp vào StarUML hoặc PlantUML extension.'
)
fmt(r, size=10, color=(30, 41, 59))

# Section 2+: Detailed 12 Use Cases
for case in ACTIVITY_CASES:
    idx = case['idx']
    title = case['title']
    uc_id = case['uc_id']
    heading_text = f"{idx}. Sơ đồ hoạt động của {title}"
    
    print(f"[{idx}/12] Generating Activity Diagram & PlantUML for {heading_text}...")
    doc_target.add_page_break()
    
    # Heading 1 (EXACT TEMPLATE REQUESTED BY USER)
    h_uc = doc_target.add_heading(level=1)
    r = h_uc.add_run(heading_text)
    fmt(r, size=13, bold=True, color=(15, 23, 42))
    
    # Overview
    p_ov = doc_target.add_paragraph()
    p_ov.paragraph_format.space_after = Pt(4)
    r = p_ov.add_run(f"Mã Use Case: {uc_id} | {case['overview']}")
    fmt(r, size=9.5, italic=True, color=(51, 65, 85))
    
    p_part = doc_target.add_paragraph()
    p_part.paragraph_format.space_after = Pt(8)
    r_lbl = p_part.add_run('Phân làn Swimlanes (Partition): ')
    fmt(r_lbl, size=9.5, bold=True, color=(30, 41, 59))
    r_val = p_part.add_run(' | '.join(case['partitions']))
    fmt(r_val, size=9.5, bold=False, color=(30, 41, 59))
    
    # Generate Activity Diagram Image
    img_filename = f"act_{idx:02d}.png"
    img_path = draw_activity_diagram(f"{uc_id} — {title}", case['partitions'], case['nodes'], img_filename)
    
    # Add Image
    doc_target.add_picture(img_path, width=Inches(6.5))
    
    # Caption
    p_cap = doc_target.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(f"Hình — {idx}: Sơ đồ hoạt động {uc_id} — {title}")
    fmt(r_cap, size=8.5, italic=True, color=(100, 116, 139))
    
    # Heading 2: Description Steps Table
    h2_tbl = doc_target.add_heading(level=2)
    r = h2_tbl.add_run('Chi tiết các bước hoạt động & Nhánh quyết định')
    fmt(r, size=11, bold=True, color=(30, 41, 59))
    
    tbl_desc = doc_target.add_table(rows=1, cols=4)
    set_table_col_widths(tbl_desc, [0.45, 1.4, 1.8, 2.85])
    headers = ['#', 'Phân làn (Partition)', 'Loại nút / Quyết định', 'Nội dung hành vi / Điều kiện']
    for c, h in zip(tbl_desc.rows[0].cells, headers):
        set_cell(c, h, bold=True, size=8.5, color=(255, 255, 255))
        shade_cell(c, '334155')
        
    step_counter = 1
    for st_part, st_type, st_text, st_cond in case['steps_def']:
        if st_type in ['action', 'decision', 'else']:
            row = tbl_desc.add_row()
            set_cell(row.cells[0], str(step_counter), bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell(row.cells[1], st_part, bold=True, size=8)
            
            type_label = 'Hành động' if st_type == 'action' else ('Nút quyết định' if st_type == 'decision' else 'Nhánh điều kiện')
            set_cell(row.cells[2], type_label, bold=False, size=8)
            
            if st_type == 'else':
                shade_cell(row.cells[2], 'FEF2F2')
                shade_cell(row.cells[3], 'FEF2F2')
                set_cell(row.cells[3], f"[{st_cond}] -> {st_text}", bold=False, size=8, color=(185, 28, 28))
            else:
                set_cell(row.cells[3], st_text, bold=False, size=8)
                
            step_counter += 1
            
    p_sp = doc_target.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)
    
    # Heading 2: PlantUML Code Block
    h2_puml = doc_target.add_heading(level=2)
    r = h2_puml.add_run(f"Mã PlantUML sơ đồ hoạt động của {title}")
    fmt(r, size=11, bold=True, color=(30, 41, 59))
    
    puml_activity = generate_activity_plantuml(title, idx, case['partitions'], case['steps_def'])
    add_code_block(doc_target, puml_activity)

# Output doc path
output_act_path = Path('Act_Diagram_1.0.docx')
try:
    doc_target.save(output_act_path)
    print('====================================================')
    print(f'SUCCESS! File saved to: {output_act_path.resolve()}')
    print('====================================================')
except PermissionError:
    alt_act_path = Path('Act_Diagram_1.0_fixed.docx')
    doc_target.save(alt_act_path)
    print('====================================================')
    print(f'Act_Diagram_1.0.docx đang mở trong Word! Đã lưu tạm sang: {alt_act_path.resolve()}')
    print('====================================================')
