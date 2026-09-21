from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import textwrap

FILE = Path('Sơ_đồ_tuần_tự.docx')
ASSETS = Path('exception_sequence_assets')
ASSETS.mkdir(exist_ok=True)

FLOWS = [
 ('EXC-001','UC-001 — Đăng nhập và phân quyền','Thông tin đăng nhập không hợp lệ',
  ['Actor','Giao diện','Dịch vụ xác thực','Kho tài khoản'],
  [('Actor','Giao diện','Gửi thông tin đăng nhập'),('Giao diện','Dịch vụ xác thực','Yêu cầu xác thực'),('Dịch vụ xác thực','Kho tài khoản','Đối chiếu tài khoản'),('Kho tài khoản','Dịch vụ xác thực','Không hợp lệ'),('Dịch vụ xác thực','Giao diện','Từ chối xác thực'),('Giao diện','Actor','Thông báo lỗi')]),
 ('EXC-002','UC-002 — Quản lý tài khoản và phân quyền','Actor không có quyền hoặc dữ liệu thay đổi không hợp lệ',
  ['Quản lý','Giao diện','Dịch vụ tài khoản','Kho tài khoản'],
  [('Quản lý','Giao diện','Gửi thay đổi'),('Giao diện','Dịch vụ tài khoản','Yêu cầu cập nhật'),('Dịch vụ tài khoản','Kho tài khoản','Kiểm tra quyền/dữ liệu'),('Kho tài khoản','Dịch vụ tài khoản','Không hợp lệ hoặc không có quyền'),('Dịch vụ tài khoản','Giao diện','Từ chối cập nhật'),('Giao diện','Quản lý','Thông báo lỗi')]),
 ('EXC-003','UC-003 — Quản lý danh mục','Dữ liệu danh mục thiếu, sai hoặc không thể lưu',
  ['Quản lý','Giao diện','Dịch vụ danh mục','Kho dữ liệu'],
  [('Quản lý','Giao diện','Gửi dữ liệu danh mục'),('Giao diện','Dịch vụ danh mục','Yêu cầu cập nhật'),('Dịch vụ danh mục','Kho dữ liệu','Kiểm tra/lưu dữ liệu'),('Kho dữ liệu','Dịch vụ danh mục','Lỗi hoặc dữ liệu không hợp lệ'),('Dịch vụ danh mục','Giao diện','Trả lỗi'),('Giao diện','Quản lý','Yêu cầu điều chỉnh')]),
 ('EXC-004','UC-004 — Ghi nhận xe vào','Không đủ dữ liệu hoặc không xác định được vị trí phù hợp',
  ['Nhân viên','Giao diện','Dịch vụ bãi xe','Kho dữ liệu'],
  [('Nhân viên','Giao diện','Gửi dữ liệu xe vào'),('Giao diện','Dịch vụ bãi xe','Yêu cầu ghi nhận'),('Dịch vụ bãi xe','Kho dữ liệu','Kiểm tra dữ liệu/tình trạng chỗ'),('Kho dữ liệu','Dịch vụ bãi xe','Thiếu dữ liệu hoặc không có vị trí'),('Dịch vụ bãi xe','Giao diện','Không thể hoàn tất'),('Giao diện','Nhân viên','Thông báo và yêu cầu điều chỉnh')]),
 ('EXC-005','UC-005 — Ghi nhận xe ra và tính phí','Không tìm thấy lượt gửi hoặc không xác định được mức phí',
  ['Nhân viên','Giao diện','Dịch vụ bãi xe','Kho dữ liệu','Bảng giá'],
  [('Nhân viên','Giao diện','Gửi thông tin xe ra'),('Giao diện','Dịch vụ bãi xe','Yêu cầu xử lý'),('Dịch vụ bãi xe','Kho dữ liệu','Tìm lượt gửi'),('Kho dữ liệu','Dịch vụ bãi xe','Không tìm thấy'),('Dịch vụ bãi xe','Bảng giá','Xác định phí nếu có lượt gửi'),('Bảng giá','Dịch vụ bãi xe','Không có mức phí hợp lệ'),('Dịch vụ bãi xe','Giao diện','Không hoàn tất lượt gửi'),('Giao diện','Nhân viên','Thông báo lỗi')]),
 ('EXC-006','UC-006 — Theo dõi chỗ trống','Không thể truy xuất dữ liệu tình trạng bãi đỗ',
  ['Actor','Giao diện','Dịch vụ theo dõi','Kho dữ liệu'],
  [('Actor','Giao diện','Yêu cầu xem chỗ trống'),('Giao diện','Dịch vụ theo dõi','Gửi yêu cầu'),('Dịch vụ theo dõi','Kho dữ liệu','Lấy tình trạng chỗ'),('Kho dữ liệu','Dịch vụ theo dõi','Không khả dụng'),('Dịch vụ theo dõi','Giao diện','Trả lỗi'),('Giao diện','Actor','Thông báo không thể hiển thị')]),
 ('EXC-007','UC-007 — Tra cứu lượt gửi xe','Không có kết quả hoặc tiêu chí tra cứu không hợp lệ',
  ['Actor','Giao diện','Dịch vụ tra cứu','Kho dữ liệu'],
  [('Actor','Giao diện','Gửi tiêu chí tra cứu'),('Giao diện','Dịch vụ tra cứu','Yêu cầu tìm kiếm'),('Dịch vụ tra cứu','Kho dữ liệu','Truy xuất lượt gửi'),('Kho dữ liệu','Dịch vụ tra cứu','Không có kết quả'),('Dịch vụ tra cứu','Giao diện','Kết quả rỗng hoặc lỗi tiêu chí'),('Giao diện','Actor','Thông báo phù hợp')]),
 ('EXC-008','UC-008 — Quản lý vé tháng/khách quen','Dữ liệu vé không hợp lệ hoặc không thể lưu',
  ['Quản lý','Giao diện','Dịch vụ vé','Kho dữ liệu'],
  [('Quản lý','Giao diện','Gửi dữ liệu vé'),('Giao diện','Dịch vụ vé','Yêu cầu cập nhật'),('Dịch vụ vé','Kho dữ liệu','Kiểm tra/lưu dữ liệu'),('Kho dữ liệu','Dịch vụ vé','Dữ liệu không hợp lệ hoặc lỗi lưu'),('Dịch vụ vé','Giao diện','Trả lỗi'),('Giao diện','Quản lý','Yêu cầu điều chỉnh')]),
 ('EXC-009','UC-009 — Xem thống kê vận hành','Không đủ dữ liệu hoặc không thể tổng hợp thống kê',
  ['Quản lý','Giao diện','Dịch vụ thống kê','Kho dữ liệu'],
  [('Quản lý','Giao diện','Yêu cầu thống kê'),('Giao diện','Dịch vụ thống kê','Gửi yêu cầu'),('Dịch vụ thống kê','Kho dữ liệu','Lấy dữ liệu lượt gửi'),('Kho dữ liệu','Dịch vụ thống kê','Không đủ dữ liệu hoặc lỗi'),('Dịch vụ thống kê','Giao diện','Không có thống kê hợp lệ'),('Giao diện','Quản lý','Thông báo phù hợp')]),
 ('EXC-010','UC-010 — AI sinh báo cáo lưu lượng','Dữ liệu không đủ, AI timeout/rate limit hoặc response không hợp lệ',
  ['Quản lý','Giao diện','Dịch vụ AI','Dịch vụ thống kê','Nhà cung cấp AI'],
  [('Quản lý','Giao diện','Yêu cầu báo cáo AI'),('Giao diện','Dịch vụ AI','Gửi yêu cầu'),('Dịch vụ AI','Dịch vụ thống kê','Lấy dữ liệu'),('Dịch vụ thống kê','Dịch vụ AI','Dữ liệu không đủ / hợp lệ'),('Dịch vụ AI','Nhà cung cấp AI','Gọi AI khi có dữ liệu'),('Nhà cung cấp AI','Dịch vụ AI','Timeout/rate limit/response sai'),('Dịch vụ AI','Giao diện','Fallback: không tạo báo cáo'),('Giao diện','Quản lý','Thông báo lỗi và yêu cầu kiểm chứng')]),
 ('EXC-011','UC-011 — AI hỏi đáp dữ liệu bãi xe','Không có dữ liệu phù hợp, dữ liệu vượt quyền hoặc AI không khả dụng',
  ['Quản lý','Giao diện','Dịch vụ AI','Dịch vụ thống kê','Nhà cung cấp AI'],
  [('Quản lý','Giao diện','Gửi câu hỏi'),('Giao diện','Dịch vụ AI','Gửi yêu cầu'),('Dịch vụ AI','Dịch vụ thống kê','Kiểm tra ngữ cảnh/quyền'),('Dịch vụ thống kê','Dịch vụ AI','Không có dữ liệu hoặc vượt quyền'),('Dịch vụ AI','Nhà cung cấp AI','Gọi AI khi hợp lệ'),('Nhà cung cấp AI','Dịch vụ AI','Không khả dụng/response sai'),('Dịch vụ AI','Giao diện','Fallback: không trả lời thay thế'),('Giao diện','Quản lý','Thông báo phù hợp')]),
 ('EXC-012','UC-012 — AI gợi ý bố trí nhân sự','Không đủ dữ liệu cao điểm hoặc AI không khả dụng',
  ['Quản lý','Giao diện','Dịch vụ AI','Dịch vụ thống kê','Nhà cung cấp AI'],
  [('Quản lý','Giao diện','Yêu cầu gợi ý'),('Giao diện','Dịch vụ AI','Gửi yêu cầu'),('Dịch vụ AI','Dịch vụ thống kê','Lấy dữ liệu cao điểm'),('Dịch vụ thống kê','Dịch vụ AI','Không đủ dữ liệu'),('Dịch vụ AI','Nhà cung cấp AI','Gọi AI khi có dữ liệu'),('Nhà cung cấp AI','Dịch vụ AI','Timeout/rate limit/response sai'),('Dịch vụ AI','Giao diện','Không đưa ra gợi ý'),('Giao diện','Quản lý','Thông báo lỗi; không tự động thay đổi vận hành')]),
]

def draw(code, condition, participants, messages):
    width=1600; height=max(650,265+len(messages)*104); image=Image.new('RGB',(width,height),'white'); d=ImageDraw.Draw(image)
    try:
        f=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24); fb=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',26); fs=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',20)
    except OSError: f=fb=fs=ImageFont.load_default()
    n=len(participants); xs=[120+i*(width-240)/(n-1) if n>1 else width/2 for i in range(n)]
    d.rounded_rectangle((35,18,width-35,62),radius=10,fill='#FFF8E8',outline='#B8860B',width=2)
    d.text((58,29),'alt — '+condition,font=f,fill='#7A5A00')
    top=110; start=195; bottom=height-52
    for x,name in zip(xs,participants):
        box=d.textbbox((0,0),name,font=fb); tw=box[2]-box[0]; th=box[3]-box[1]
        d.rounded_rectangle((x-tw/2-22,top-th/2-16,x+tw/2+22,top+th/2+16),radius=11,fill='#F4F6F9',outline='#9B1C1C',width=3)
        d.text((x-tw/2,top-th/2),name,font=fb,fill='#5A1111')
        for y in range(start,bottom,18): d.line((x,y,x,min(y+9,bottom)),fill='#A98E8E',width=2)
    for i,(src,dst,msg) in enumerate(messages):
        y=start+i*82; a=xs[participants.index(src)]; b=xs[participants.index(dst)]
        d.line((a,y,b,y),fill='#9B1C1C',width=3); direction=1 if b>a else -1; d.polygon([(b,y),(b-direction*18,y-9),(b-direction*18,y+9)],fill='#9B1C1C')
        text='\n'.join(textwrap.wrap(msg,width=26)); bb=d.multiline_textbbox((0,0),text,font=fs,spacing=2); tw=bb[2]-bb[0]; th=bb[3]-bb[1]; lx=(a+b)/2-tw/2; ly=y-th-9
        d.rectangle((lx-6,ly-3,lx+tw+6,ly+th+3),fill='white'); d.multiline_text((lx,ly),text,font=fs,fill='#5A1111',spacing=2,align='center')
    d.text((30,height-38),code,font=fs,fill='#7A5A00'); path=ASSETS/f'{code}.png'; image.save(path); return path

def fmt(run,size=10,bold=None,color=None,name='Calibri'):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name); run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor(*color)

def set_cell(cell,text,bold=False,size=8.5):
    p=cell.paragraphs[0]; p.clear(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05; r=p.add_run(text); fmt(r,size,bold); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def widths(table,vals):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell,w in zip(row.cells,vals): cell.width=Inches(w)

def code_block(doc,code):
    t=doc.add_table(rows=1,cols=1); widths(t,[6.5]); c=t.cell(0,0); shade(c,'FFF8E8'); p=c.paragraphs[0]; p.clear(); p.paragraph_format.space_after=Pt(0); r=p.add_run(code); fmt(r,7.3,False,(122,90,0),'Consolas')

def plantuml(code,title,condition,participants,messages):
    out=['@startuml',f'title {code} — {title}','hide footbox','skinparam sequenceMessageAlign center']; aliases={}
    for i,p in enumerate(participants): aliases[p]=f'P{i+1}'; out.append(f'participant "{p}" as P{i+1}')
    out.append(f'alt {condition}')
    for s,d,m in messages: out.append(f'{aliases[s]} -> {aliases[d]} : {m}')
    out.extend(['else Luồng chính','  note over P1: Tiếp tục luồng chính khi điều kiện ngoại lệ không xảy ra','end','@enduml']); return '\n'.join(out)

doc=Document(FILE)
doc.add_page_break(); doc.add_heading('Luồng ngoại lệ bổ sung',level=1)
doc.add_paragraph('Các sơ đồ dưới đây bổ sung nhánh `alt`/`else` cho từng Use Case. Nhánh ngoại lệ chỉ mô tả xử lý lỗi hoặc điều kiện không thể hoàn tất luồng chính; không tự tạo nghiệp vụ mới.')
for code,title,condition,participants,messages in FLOWS:
    doc.add_page_break(); doc.add_heading(f'{code} — {title}',level=1)
    p=doc.add_paragraph(); r=p.add_run('Điều kiện ngoại lệ: '); fmt(r,11,True,(122,90,0)); r=p.add_run(condition); fmt(r,11,False,(122,90,0))
    doc.add_picture(str(draw(code,condition,participants,messages)),width=Inches(6.45))
    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=cap.add_run(f'Hình — {code}: luồng ngoại lệ'); fmt(r,9,False,(122,90,0))
    doc.add_heading('Thông điệp ngoại lệ',level=2)
    t=doc.add_table(rows=1,cols=4); widths(t,[0.55,1.35,1.35,3.25])
    for c,x in zip(t.rows[0].cells,['#','Từ','Đến','Thông điệp']): set_cell(c,x,True); shade(c,'FFF8E8')
    for i,(s,d,m) in enumerate(messages,1):
        cells=t.add_row().cells
        for c,x in zip(cells,[str(i),s,d,m]): set_cell(c,x)
    doc.add_heading('Mã PlantUML với nhánh alt/else',level=2); code_block(doc,plantuml(code,title,condition,participants,messages))
doc.save(FILE)
print(FILE.resolve())
